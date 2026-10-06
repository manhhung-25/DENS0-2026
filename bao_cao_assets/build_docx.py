"""Build an editable academic Word report from BAO_CAO_DO_AN.md.

Run with the Codex document runtime's Python, which provides python-docx and Pillow.
"""

from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "BAO_CAO_DO_AN.md"
OUTPUT = ROOT / "BAO_CAO_DO_AN.docx"
DIAGRAM = ROOT / "bao_cao_assets" / "hinh_3_kien_truc_word.png"
INK = RGBColor(0, 0, 0)
GREEN = RGBColor(19, 78, 51)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "5")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)


def set_cell_margin(cell, top=105, start=110, bottom=105, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        mar.append(node)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_font(run, name="Times New Roman", size=None, color=INK) -> None:
    run.font.name = name
    run.font.color.rgb = color
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is not None:
        for tag in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{tag}"), name)


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^]]+\]\([^)]+\)|<br>)")


def add_inline(paragraph, value: str, *, size=None, default_bold=False) -> None:
    parts = INLINE.split(value)
    for part in parts:
        if not part:
            continue
        if part == "<br>":
            paragraph.add_run().add_break()
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Consolas", 9 if size is None else size - 2)
        elif part.startswith("[") and "](" in part:
            label, url = part[1:-1].split("](", 1)
            add_hyperlink(paragraph, label, url)
            continue
        else:
            run = paragraph.add_run(part)
            run.bold = default_bold
        if not (part.startswith("`") and part.endswith("`")):
            set_font(run, size=size)


def add_hyperlink(paragraph, label: str, url: str) -> None:
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rpr.append(rfonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "145C3A")
    rpr.append(color)
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run._r.append(instr)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    run = paragraph.add_run("Cập nhật mục lục trong Word: bấm chuột phải → Update Field")
    set_font(run, size=10)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("Trang ")
    for kind, content in (("begin", None), ("instruction", "PAGE"), ("end", None)):
        run = paragraph.add_run()
        if kind == "instruction":
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = content
        else:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        run._r.append(element)


def add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    ncols = max(map(len, rows))
    table = doc.add_table(rows=0, cols=ncols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    full = 16.0
    if ncols == 2:
        widths = [3.5, 12.5]
    elif ncols == 3:
        widths = [3.3, 5.1, 7.6]
    elif ncols == 4:
        widths = [3.0, 3.1, 4.8, 5.1]
    elif ncols == 5:
        widths = [2.7, 1.6, 4.4, 2.0, 5.3]
    else:
        widths = [full / ncols] * ncols
    for ridx, row_values in enumerate(rows):
        row = table.add_row()
        if ridx == 0:
            set_repeat_header(row)
        for cidx, cell in enumerate(row.cells):
            cell.width = Cm(widths[cidx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margin(cell)
            if ridx == 0:
                set_cell_shading(cell, "164C35")
            elif ridx % 2 == 0:
                set_cell_shading(cell, "F5F7F5")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if cidx < len(row_values):
                add_inline(p, row_values[cidx].strip(), size=9.5, default_bold=(ridx == 0))
            if ridx == 0:
                for run in p.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            if cidx == 0 and len(row_values[cidx]) < 22:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def draw_diagram() -> None:
    image = Image.new("RGB", (1700, 1040), "white")
    draw = ImageDraw.Draw(image)
    font_path = r"C:\Windows\Fonts\times.ttf"
    bold_path = r"C:\Windows\Fonts\timesbd.ttf"
    title_font = ImageFont.truetype(bold_path, 37)
    body_font = ImageFont.truetype(font_path, 32)
    small_font = ImageFont.truetype(font_path, 28)

    def box(xy, lines, fill="#f3f7f3", font=body_font):
        draw.rounded_rectangle(xy, radius=18, fill=fill, outline="#145c3a", width=3)
        x0, y0, x1, y1 = xy
        spacing = 11
        total = len(lines) * (font.size + spacing) - spacing
        y = y0 + ((y1 - y0) - total) / 2
        for line in lines:
            bbox = draw.textbbox((0, 0), line, font=font)
            width = bbox[2] - bbox[0]
            draw.text((x0 + (x1 - x0 - width) / 2, y), line, font=font, fill="#101010")
            y += font.size + spacing

    draw.text((70, 28), "Kiến trúc giám sát và xác nhận bảo trì", font=title_font, fill="#101010")
    for idx, lines in enumerate((["Camera RGB"], ["Controller", "encoder"], ["Rung  ·  âm  ·  nhiệt"])):
        box((70 + idx * 550, 105, 500 + idx * 550, 235), lines)
    box((290, 330, 1410, 445), ["Gateway và đồng hồ chung  →  Kiểm tra chất lượng, căn chỉnh thời gian"])
    box((70, 535, 810, 665), ["Ước lượng pose", "và đặc trưng theo khớp"])
    box((890, 535, 1630, 665), ["Baseline theo recipe,", "tải và tốc độ"])
    box((290, 755, 1410, 865), ["Hợp nhất bằng chứng  →  Cảnh báo  →  Dashboard"])
    box((290, 900, 1410, 1010), ["Kỹ thuật viên xác nhận  →  Hồ sơ bảo trì  →  Đánh giá lại mô hình"], font=small_font)
    for x in (285, 835, 1385):
        draw.line((x, 235, x, 310), fill="#145c3a", width=5)
    draw.line((850, 445, 850, 520), fill="#145c3a", width=5)
    for x in (440, 1260):
        draw.line((x, 665, x, 740), fill="#145c3a", width=5)
    draw.line((850, 865, 850, 890), fill="#145c3a", width=5)
    image.save(DIAGRAM)


def make_docx() -> None:
    draw_diagram()
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.header_distance = Cm(1.0)
    sec.footer_distance = Cm(1.1)
    sec.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = INK
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for name, size, before, after in (("Title", 20, 16, 10), ("Heading 1", 15, 15, 10), ("Heading 2", 13, 12, 7), ("Heading 3", 12, 10, 5)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = INK
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        rpr = style._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is not None:
            for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                rfonts.attrib.pop(qn(f"w:{attr}"), None)
            for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
                rfonts.set(qn(f"w:{attr}"), "Times New Roman")
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    old_border = title_ppr.find(qn("w:pBdr"))
    if old_border is not None:
        title_ppr.remove(old_border)
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "nil")
    border.append(bottom)
    title_ppr.append(border)
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings.element.append(update_fields)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    in_toc = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if i == 0 and line == "# BÁO CÁO ĐỒ ÁN":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(80)
            r = p.add_run("BÁO CÁO ĐỒ ÁN")
            set_font(r, size=14)
            r.bold = True
            i += 1
            continue
        if line.startswith("# Thiết kế hệ thống"):
            doc.add_paragraph(line[2:], style="Title")
            i += 1
            continue
        if line == "## MỤC LỤC":
            doc.add_page_break()
            doc.add_paragraph("MỤC LỤC", style="Heading 1")
            p = doc.add_paragraph()
            add_field(p, ' TOC \\o "1-2" \\h \\z \\u ')
            in_toc = True
            i += 1
            continue
        if in_toc:
            if line.startswith("## PHẦN MỞ ĐẦU"):
                in_toc = False
            else:
                i += 1
                continue
        if line in ("## TÓM TẮT", "## LỜI CAM ĐOAN", "## DANH MỤC CHỮ VIẾT TẮT", "## DANH MỤC BẢNG", "## PHẦN MỞ ĐẦU", "## PHẦN NỘI DUNG"):
            doc.add_page_break()
        if line.startswith("# CHƯƠNG") or line.startswith("# KẾT LUẬN") or line.startswith("# DANH MỤC TÀI LIỆU") or line.startswith("# PHỤ LỤC"):
            doc.add_page_break()
        if line.startswith("# "):
            doc.add_paragraph(line[2:], style="Heading 1")
            i += 1
            continue
        if line.startswith("## "):
            text = line[3:]
            style = "Heading 2" if re.match(r"\d+\.\d+\.", text) else "Heading 1"
            doc.add_paragraph(text, style=style)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
            i += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip()
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            if language == "mermaid":
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.add_run().add_picture(str(DIAGRAM), width=Cm(15.8))
            else:
                for code_line in code:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.left_indent = Cm(0.5)
                    r = p.add_run(code_line or " ")
                    set_font(r, "Consolas", 8.5)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        if line.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                values = [value.strip() for value in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", value) for value in values):
                    rows.append(values)
                i += 1
            add_table(doc, rows)
            continue
        image_match = re.match(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image_match:
            path = ROOT / image_match.group(2)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(path), width=Cm(15.5))
            i += 1
            continue
        if line.startswith("**Bảng") or line.startswith("**Hình"):
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(6)
            add_inline(p, line, size=10.5)
            i += 1
            continue
        if line.startswith("> "):
            if "Các trang trong sườn mẫu" in line:
                i += 1
                continue
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, line[2:], size=10.5)
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", line))
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            i += 1
            continue
        paragraph_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(("#", "|", "```", "![", "> ", "- ")):
            if re.match(r"^\d+\. ", lines[i].strip()):
                break
            paragraph_lines.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines))

    doc.core_properties.title = "Thiết kế hệ thống giám sát tư thế robot Franka Panda và mô phỏng cảnh báo bảo trì đa cảm biến"
    doc.core_properties.subject = "Báo cáo đồ án robot và bảo trì đa cảm biến"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    make_docx()
